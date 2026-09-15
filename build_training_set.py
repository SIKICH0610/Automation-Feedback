#!/usr/bin/env python3
"""Build the LoRA training set from the labeled corpus spreadsheet.

Reads ~/Downloads/语料汇总-labeling.xlsx, keeps rows the teacher marked 好,
applies the sentence fixes the teacher asked for in 备注, strips greetings and
student names, and writes:

    app_data/training/pairs_v1.jsonl    all approved 素材→评语 pairs (zh + en)
    app_data/training/singles_v1.jsonl  approved comments without 素材 (style bank)
    app_data/training/train.jsonl       MLX-LM chat format, Chinese pairs
    app_data/training/valid.jsonl       held-out validation split
    ~/Downloads/训练集v1-入选评语.xlsx    human-readable review copy

Labels typed into the sheet win; LABEL_PATCH fills rows whose 标注 cell is
still empty (the first labeling round arrived via chat before the sheet was
saved). Re-run after more labeling:  .venv/bin/python build_training_set.py
"""
from __future__ import annotations

import json
import random
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

LABEL_XLSX = Path.home() / "Downloads" / "语料汇总-labeling.xlsx"
OUT_DIR = Path(__file__).resolve().parent / "app_data" / "training"
REVIEW_XLSX = Path.home() / "Downloads" / "训练集v1-入选评语.xlsx"
VALID_SIZE = 4
SEED = 42

# First labeling round (2026-09-15), transcribed from the teacher's pasted
# sheet; a non-empty 标注 cell in the file always overrides this.
_GOOD = ("好", "")
LABEL_PATCH: dict[int, tuple[str, str]] = {
    1: _GOOD,
    2: ("好", "有ai味"),
    3: ("不好", ""),
    4: ("待定", ""),
    5: ("不好", "有ai味"),
    6: ("不好", "有ai味"),
    7: _GOOD,
    8: _GOOD,
    9: ("不好", "不适合所有课程"),
    10: ("待定", ""),
    11: ("不好", "不适合所有课程"),
    12: ("待定", ""),
    13: ("待定", ""),
    14: ("好", "语气可以更友善"),
    15: ("好", "珍贵的多有点奇怪，可以改成“好习惯”"),
    16: ("好", "有ai味"),
    17: ("好", "有ai味"),
    18: _GOOD,
    19: _GOOD,
    20: _GOOD,
    21: _GOOD,
    22: _GOOD,
    23: ("好", "有ai味"),
    24: ("好", "有ai味，“这股劲我想帮她保住”很奇怪"),
    25: ("好", "有ai味，“所以安静不是没听进去，是性格。”很奇怪"),
    26: _GOOD,
    27: _GOOD,
    28: _GOOD,
    29: _GOOD,
    30: _GOOD,
    31: _GOOD,
    32: ("好", "有ai味"),
    33: _GOOD,
    34: ("好", "有ai味"),
    **{i: _GOOD for i in range(35, 49)},
}

# Sentence-level repairs the teacher requested in 备注, applied to the 评语
# before it enters the corpus. Keyed by 编号.
FIXUPS: dict[int, list[tuple[str, str]]] = {
    14: [
        ("，她有上了3D图形的课程不", "。"),
        ("，有任何问题随时找我～", "。"),
    ],
    15: [("这个习惯比多做对几道题珍贵得多", "这是非常好的习惯")],
    24: [("这股劲儿我想帮她保住", "希望她能把这股劲头保持下去")],
    25: [("所以安静不是没听进去，是性格。", "她只是性格偏安静，内容其实都听进去了。")],
}
# Full-message rows: keep only the personal paragraph, from this substring on
# (the recap/greeting before it is the template's job, not the model's).
TRIM_TO: dict[int, str] = {
    14: "Joelle是第二节课才来的",
}

_GREETING = re.compile(r"^\s*(?:[A-Za-z][\w .'-]*)?家长您好[，,。！!～~\s]*")

TRAIN_SYSTEM = (
    "你是 Think Academy 的数学老师，正在给家长写孩子的课堂反馈。"
    "把老师的课堂速记/关键词改写成发给家长的评语：保留速记里的每个事实，不添加新事实；"
    "不写学生姓名和称呼，不写问候和结尾；这条消息只发给一位家长、只谈这一个孩子，"
    "用“孩子”做主语，绝不能出现“孩子们”“同学们”等群体称呼。"
)


def _has_cjk(text: str) -> bool:
    return any("一" <= ch <= "鿿" for ch in text)


def _strip_name(text: str, student: str, lang: str) -> str:
    """Replace the student's name so the corpus teaches voice, not names."""
    result = _GREETING.sub("", text.strip())
    for token in (student or "").split():
        if len(token) < 3:
            continue  # too short to replace safely ("Q", "Me", "An")
        placeholder = "孩子" if lang == "zh" else "the student"
        # \b fails against CJK neighbours (they count as \w), so bound the
        # name against Latin letters only: "Katherine这节课" must still match.
        result = re.sub(
            rf"(?<![A-Za-z]){re.escape(token)}(?![A-Za-z])",
            placeholder,
            result,
            flags=re.IGNORECASE,
        )
    if lang == "zh":
        result = result.replace("孩子孩子", "孩子")
    else:
        # Capitalize placeholders that start a sentence.
        result = re.sub(
            r"(^|[.!?]\s+)the student", lambda m: m.group(1) + "The student", result
        )
    return result.strip()


def load_rows() -> list[dict]:
    wb = load_workbook(LABEL_XLSX, data_only=True)
    ws = wb.active
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        rid, rtype, material, final, _count, student, source, label, note = (
            (list(raw) + [None] * 9)[:9]
        )
        if rid is None or not final:
            continue
        rid = int(rid)
        label = str(label).strip() if label else ""
        note = str(note).strip() if note else ""
        if not label and rid in LABEL_PATCH:
            label, note = LABEL_PATCH[rid]
        rows.append(
            {
                "id": rid,
                "type": str(rtype or "").strip(),
                "material": str(material or "").strip(),
                "final": str(final or "").strip(),
                "student": str(student or "").strip(),
                "source": str(source or "").strip(),
                "label": label,
                "note": note,
            }
        )
    return rows


def build() -> None:
    rows = load_rows()
    labeled = [r for r in rows if r["label"]]
    good = [r for r in labeled if r["label"] == "好"]

    pairs, singles = [], []
    for row in good:
        final = row["final"]
        if row["id"] in TRIM_TO:
            anchor = TRIM_TO[row["id"]]
            if anchor in final:
                final = final[final.index(anchor):]
        for old, new in FIXUPS.get(row["id"], []):
            final = final.replace(old, new)
        lang = "zh" if _has_cjk(final) else "en"
        entry = {
            "id": row["id"],
            "quality": "B" if row["note"] else "A",
            "lang": lang,
            "source": row["source"],
            "material": _strip_name(row["material"], row["student"], lang),
            "final": _strip_name(final, row["student"], lang),
            "note": row["note"],
        }
        (pairs if entry["material"] else singles).append(entry)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, data in (("pairs_v1.jsonl", pairs), ("singles_v1.jsonl", singles)):
        with open(OUT_DIR / name, "w", encoding="utf-8") as fh:
            for entry in data:
                fh.write(json.dumps(entry, ensure_ascii=False) + "\n")

    zh_pairs = [p for p in pairs if p["lang"] == "zh"]
    shuffled = zh_pairs[:]
    random.Random(SEED).shuffle(shuffled)
    valid, train = shuffled[:VALID_SIZE], shuffled[VALID_SIZE:]
    for name, split in (("train.jsonl", train), ("valid.jsonl", valid)):
        with open(OUT_DIR / name, "w", encoding="utf-8") as fh:
            for entry in split:
                sample = {
                    "messages": [
                        {"role": "system", "content": TRAIN_SYSTEM},
                        {"role": "user", "content": f"课堂速记：{entry['material']}"},
                        {"role": "assistant", "content": entry["final"]},
                    ]
                }
                fh.write(json.dumps(sample, ensure_ascii=False) + "\n")

    write_review(pairs + singles)

    total = len(rows)
    print(f"sheet rows: {total}, labeled: {len(labeled)}, 好: {len(good)}")
    print(f"pairs: {len(pairs)} (zh {len(zh_pairs)}, en {len(pairs) - len(zh_pairs)}), singles: {len(singles)}")
    print(f"quality A: {sum(1 for p in pairs if p['quality'] == 'A')}, B: {sum(1 for p in pairs if p['quality'] == 'B')}")
    print(f"train: {len(train)}, valid: {len(valid)} -> {OUT_DIR}")
    print(f"review copy -> {REVIEW_XLSX}")


def write_review(entries: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "训练集v1"
    headers = ["原表编号", "质量", "语言", "素材（清洗后）", "评语（入选定稿）", "来源", "处理说明"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for entry in sorted(entries, key=lambda e: e["id"]):
        notes = []
        if entry["id"] in FIXUPS or entry["id"] in TRIM_TO:
            notes.append("已按你的备注修改")
        if entry["note"]:
            notes.append(f"标注备注：{entry['note']}")
        ws.append(
            [
                entry["id"],
                entry["quality"],
                entry["lang"],
                entry["material"],
                entry["final"],
                entry["source"],
                "；".join(notes),
            ]
        )
    widths = {"A": 9, "B": 6, "C": 6, "D": 46, "E": 64, "F": 26, "G": 34}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(REVIEW_XLSX)


if __name__ == "__main__":
    build()
