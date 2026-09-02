from __future__ import annotations

import argparse
import codecs
import csv
import io
import json
import warnings
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from database_store import SQLiteFeedbackStore
from feedback_common import DEFAULT_WORKBOOK


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_APP_DATA_DIR = PROJECT_DIR / "app_data"
DEFAULT_DATABASE = DEFAULT_APP_DATA_DIR / "feedback.db"
DEFAULT_EXPORT_DIR = PROJECT_DIR / "exports"
DEFAULT_ANNOUNCEMENT_DIR = PROJECT_DIR / "announcements"

# Canonical field -> header spellings accepted for it, in normalized form (see
# _normalize_header). The canonical names are what group_by_class reads. Matching
# is by name and alias rather than by position, extra columns are ignored, and the
# header row is located by content -- so a platform export that adds columns,
# reorders them, renames them within these aliases, or grows a banner row above
# the table imports without any code change.
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "classId": ("classid", "班级id", "课程id"),
    "className": ("classname", "班级名称", "课程名称", "班级名"),
    "classTimeDescription": ("classtimedescription", "classtime", "上课时间"),
    "subject": ("subject", "科目", "学科"),
    "firstName": ("firstname", "givenname", "名"),
    "lastName": ("lastname", "familyname", "surname", "姓"),
    "学员id": ("学员id", "studentid", "学生id", "学员编号"),
    "payStatus": ("paystatus", "支付状态", "付款状态", "缴费状态"),
    "是否入班": ("是否入班",),
    "In group": ("ingroup", "是否入群"),
}

# Without these the file cannot be imported at all. The class time is required on
# purpose: it distinguishes same-named classes and teachers rely on it, so a file
# that lost it should fail loudly rather than import silently without times.
# subject and In group are cosmetic: when missing they come back empty.
REQUIRED_FIELDS = (
    "classId",
    "className",
    "classTimeDescription",
    "firstName",
    "lastName",
    "学员id",
    "payStatus",
    "是否入班",
)

# How many top rows to scan for the header row.
HEADER_SCAN_ROWS = 10


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    for junk in (" ", "\u3000", "_", "-"):
        text = text.replace(junk, "")
    return text


_ALIAS_TO_FIELD = {
    alias: field for field, aliases in FIELD_ALIASES.items() for alias in aliases
}


def _headers_in_row(cells: list[Any]) -> tuple[dict[str, int], dict[str, int]]:
    """(canonical field -> column index, raw header -> column index) for one row."""
    fields: dict[str, int] = {}
    raw: dict[str, int] = {}
    for column, value in enumerate(cells):
        text = str(value or "").strip()
        if not text:
            continue
        raw.setdefault(text, column)
        field = _ALIAS_TO_FIELD.get(_normalize_header(text))
        if field and field not in fields:
            fields[field] = column
    return fields, raw


# Strict OOXML uses these namespace roots instead of the transitional ones.
# openpyxl silently loads such a file as zero worksheets; the platform's exporter
# produces exactly that ("conformance=strict", purl.oclc.org namespaces).
_STRICT_TO_TRANSITIONAL = (
    (
        b"http://purl.oclc.org/ooxml/spreadsheetml/main",
        b"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    ),
    (
        b"http://purl.oclc.org/ooxml/officeDocument/relationships",
        b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    ),
)


def _load_workbook_lenient(xlsx_bytes: bytes):
    """load_workbook from bytes, plus a fallback for Strict OOXML files.

    Loading from bytes (not a path) also skips openpyxl's extension check, so a
    workbook that arrived under the wrong file name still opens. For Strict
    OOXML -- which openpyxl silently loads as zero worksheets -- the two dialects
    are structurally identical for our purposes; rewriting the namespace URIs to
    the transitional ones is enough to read the data. Ordinary files pay nothing.
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        if workbook.worksheets:
            return workbook
        workbook.close()

        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as source, zipfile.ZipFile(
            buffer, "w", zipfile.ZIP_DEFLATED
        ) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename.endswith((".xml", ".rels")):
                    for strict, transitional in _STRICT_TO_TRANSITIONAL:
                        data = data.replace(strict, transitional)
                target.writestr(item, data)
        buffer.seek(0)
        return load_workbook(buffer, data_only=True)


def _sniff_format(data: bytes) -> str:
    """xlsx / json / csv, decided by content so a misnamed file still imports."""
    if data.startswith(b"PK\x03\x04"):
        return "xlsx"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise ValueError(
            "This is a legacy .xls file. Re-export it as .xlsx or .csv and try again."
        )
    head = data[:4096]
    if head.startswith(codecs.BOM_UTF8):
        head = head[len(codecs.BOM_UTF8):]
    stripped = head.lstrip()
    if stripped[:1] in (b"{", b"["):
        return "json"
    return "csv"


def _decode_text(data: bytes) -> str:
    # Platform exports are UTF-8 (often with BOM); Excel-saved Chinese CSVs are GBK.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030")


def _grid_from_csv(data: bytes) -> list[list[Any]]:
    text = _decode_text(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _grids_from_xlsx(xlsx_bytes: bytes) -> list[list[list[Any]]]:
    workbook = _load_workbook_lenient(xlsx_bytes)
    try:
        return [
            [list(row) for row in worksheet.iter_rows(values_only=True)]
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _rows_from_grids(grids: list[list[list[Any]]], source_name: str) -> list[dict[str, Any]]:
    best: tuple[int, list[list[Any]], int, dict[str, int], dict[str, int]] | None = None
    for grid in grids:
        for row_index in range(min(HEADER_SCAN_ROWS, len(grid))):
            fields, raw = _headers_in_row(grid[row_index])
            found = sum(1 for field in REQUIRED_FIELDS if field in fields)
            if best is None or found > best[0]:
                best = (found, grid, row_index, fields, raw)
            if found == len(REQUIRED_FIELDS):
                break
        else:
            continue
        break

    if best is None or best[0] < len(REQUIRED_FIELDS):
        fields = best[3] if best else {}
        raw = best[4] if best else {}
        missing = [field for field in REQUIRED_FIELDS if field not in fields]
        seen = ", ".join(sorted(raw)) or "(no headers found)"
        raise ValueError(
            f"{source_name} has no row containing the required column(s): "
            f"{', '.join(missing)}. Closest header row contained: {seen}"
        )

    _, grid, header_index, fields, raw = best
    # Canonical fields win their columns; every other column rides along under
    # its own header so future optional lookups keep working.
    taken = set(fields.values())
    columns = dict(fields)
    for name, column in raw.items():
        if column not in taken and name not in columns:
            columns[name] = column

    rows: list[dict[str, Any]] = []
    for cells in grid[header_index + 1:]:
        def cell(column: int) -> Any:
            return cells[column] if column < len(cells) else None

        first_name = cell(columns["firstName"])
        if first_name is None or not str(first_name).strip():
            continue
        rows.append({name: cell(column) for name, column in columns.items()})
    return rows


def _rows_from_json(data: bytes, source_name: str) -> list[dict[str, Any]]:
    try:
        payload = json.loads(_decode_text(data))
    except json.JSONDecodeError as exc:
        raise ValueError(f"{source_name} is not valid JSON: {exc}") from exc

    records = None
    if isinstance(payload, list):
        records = payload
    elif isinstance(payload, dict):
        # Platform APIs wrap the list: {"data": [...]}, {"data": {"list": [...]}} etc.
        queue = list(payload.values())
        depth_left = len(queue) * 8
        while queue and depth_left > 0:
            depth_left -= 1
            value = queue.pop(0)
            if isinstance(value, list) and value and all(isinstance(item, dict) for item in value):
                records = value
                break
            if isinstance(value, dict):
                queue.extend(value.values())
    if records is None:
        raise ValueError(f"{source_name} does not contain a list of student records.")

    rows: list[dict[str, Any]] = []
    seen_fields: set[str] = set()
    seen_keys: set[str] = set()
    for record in records:
        if not isinstance(record, dict):
            continue
        row: dict[str, Any] = {}
        for key, value in record.items():
            text = str(key or "").strip()
            if not text:
                continue
            seen_keys.add(text)
            field = _ALIAS_TO_FIELD.get(_normalize_header(text))
            if field:
                row.setdefault(field, value)
                seen_fields.add(field)
            else:
                row.setdefault(text, value)
        first_name = str(row.get("firstName") or "").strip()
        if first_name:
            rows.append(row)

    missing = [field for field in REQUIRED_FIELDS if field not in seen_fields]
    if missing:
        seen = ", ".join(sorted(seen_keys)) or "(no keys found)"
        raise ValueError(
            f"{source_name} is missing the required field(s): "
            f"{', '.join(missing)}. Records contained: {seen}"
        )
    return rows


def normalize_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_enrollment_rows(xlsx_path: Path) -> list[dict[str, Any]]:
    """Read an enrollment export: .xlsx (transitional or strict), .csv, or .json.

    The format is sniffed from the file's bytes, not its name, so an upload that
    lost its extension -- or carries the wrong one -- still imports.
    """
    source = Path(xlsx_path)
    data = source.read_bytes()
    file_format = _sniff_format(data)
    if file_format == "json":
        return _rows_from_json(data, source.name)
    if file_format == "csv":
        grids = [_grid_from_csv(data)]
    else:
        grids = _grids_from_xlsx(data)
    return _rows_from_grids(grids, source.name)


def group_by_class(
    rows: list[dict[str, Any]],
    *,
    class_ids: set[str] | None,
) -> dict[str, dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    seen_uids_by_class: dict[str, set[str]] = {}

    for row in rows:
        if str(row.get("payStatus") or "").strip().lower() != "paid":
            continue
        if str(row.get("是否入班") or "").strip() != "是":
            continue

        class_id = normalize_id(row.get("classId"))
        if not class_id:
            continue
        if class_ids and class_id not in class_ids:
            continue

        uid = normalize_id(row.get("学员id"))
        if not uid:
            continue

        seen_uids = seen_uids_by_class.setdefault(class_id, set())
        if uid in seen_uids:
            continue
        seen_uids.add(uid)

        group = groups.setdefault(
            class_id,
            {
                "external_class_id": class_id,
                "name": str(row.get("className") or "").strip(),
                "weekly_time": str(row.get("classTimeDescription") or "").strip(),
                "subject": str(row.get("subject") or "").strip(),
                "students": [],
            },
        )
        in_group_value = row.get("In group")
        group["students"].append(
            {
                "uid": uid,
                "first_name": str(row.get("firstName") or "").strip(),
                "last_name": str(row.get("lastName") or "").strip(),
                "in_group": bool(in_group_value) if in_group_value is not None else None,
            }
        )

    return groups


def print_plan(plan: dict[str, Any]) -> None:
    semester_state = "existing" if plan["semester_exists"] else "will be created"
    print(f"Semester: {plan['semester']!r} ({semester_state})")

    for class_plan in plan["classes"]:
        class_state = "existing class" if class_plan["class_exists"] else "NEW class"
        print()
        print(f"[{class_state}] {class_plan['name']} (classId={class_plan['external_class_id']})")
        print(f"    weekly time: {class_plan['weekly_time'] or '(none)'}")
        new_students = [s for s in class_plan["students"] if not s["already_present"]]
        already_present = [s for s in class_plan["students"] if s["already_present"]]
        print(f"    {len(new_students)} new student(s), {len(already_present)} already on roster")
        for student in new_students:
            print(f"      + {student['first_name']} {student['last_name']} (uid={student['uid']})")
        to_remove = class_plan.get("students_to_remove") or []
        if to_remove:
            print(f"    {len(to_remove)} student(s) WILL BE PERMANENTLY REMOVED (not in the uploaded file):")
            for student in to_remove:
                print(f"      - {student['first_name']} {student['last_name']} (uid={student['uid']})")


def total_students_to_remove(plan: dict[str, Any]) -> int:
    return sum(len(class_plan.get("students_to_remove") or []) for class_plan in plan["classes"])


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Import a platform enrollment export into the SQLite roster, grouped into a "
            "named semester block. Prints a preview by default; pass --commit to write."
        )
    )
    parser.add_argument("--source-file", type=Path, required=True, help="Enrollment export (.xlsx, .csv, or .json).")
    parser.add_argument(
        "--semester",
        required=True,
        help="Semester block name, e.g. '2026-27 School Year'. Created if it doesn't exist yet, "
        "reused otherwise.",
    )
    parser.add_argument(
        "--class-id",
        action="append",
        default=[],
        help="Only import this platform classId. Repeat for multiple classIds. Omit to import "
        "every paid, enrolled class found in the file.",
    )
    parser.add_argument(
        "--commit",
        action="store_true",
        help="Write the changes. Without this flag, only a preview is printed and nothing is saved.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help=(
            "Sync each matched class's roster to exactly match the uploaded file: a student "
            "currently on the roster but missing from the file is permanently deleted. A "
            "student present in both is refreshed (name, Group Chat) but keeps all their other "
            "data (feedback, quiz scores, notes). Without this flag, import only ever adds new "
            "students and never touches or removes existing ones."
        ),
    )
    parser.add_argument(
        "--confirm-delete",
        action="store_true",
        help="Required together with --overwrite --commit whenever the preview shows students "
        "that would be removed, as an explicit acknowledgment of permanent deletion.",
    )
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--app-data-dir", type=Path, default=DEFAULT_APP_DATA_DIR)
    parser.add_argument("--export-dir", type=Path, default=DEFAULT_EXPORT_DIR)
    parser.add_argument("--announcement-dir", type=Path, default=DEFAULT_ANNOUNCEMENT_DIR)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    rows = read_enrollment_rows(args.source_file)
    class_ids = {normalize_id(value) for value in args.class_id} if args.class_id else None
    groups = group_by_class(rows, class_ids=class_ids)

    if not groups:
        print("No paid, enrolled rows matched the given filters. Nothing to import.")
        return

    store = SQLiteFeedbackStore(
        database_path=args.database,
        source_workbook_path=args.workbook,
        announcement_dir=args.announcement_dir,
        app_data_dir=args.app_data_dir,
        export_dir=args.export_dir,
    )
    if args.overwrite and args.commit and not args.confirm_delete:
        preview_plan = store.import_students(
            semester_name=args.semester,
            class_groups=list(groups.values()),
            commit=False,
            overwrite=True,
        )
        removal_count = total_students_to_remove(preview_plan)
        if removal_count:
            print_plan(preview_plan)
            print()
            raise SystemExit(
                f"--overwrite would permanently remove {removal_count} student(s), shown above. "
                "Re-run with --confirm-delete added to actually do this."
            )

    plan = store.import_students(
        semester_name=args.semester,
        class_groups=list(groups.values()),
        commit=args.commit,
        overwrite=args.overwrite,
    )
    print_plan(plan)
    print()
    if args.commit:
        print("Committed to the database.")
        store.ensure_announcement_files()
    else:
        print("Dry run only. Nothing was written. Re-run with --commit to save these changes.")


if __name__ == "__main__":
    main()
